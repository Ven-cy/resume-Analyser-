
from pathlib import Path
import json

import joblib
import pandas as pd
import streamlit as st

from docx import Document
from pypdf import PdfReader

from matcher import (
    analyze_match,
    find_similar_resumes,
)


BASE_DIR = Path(__file__).resolve().parent
MODEL_DIR = BASE_DIR / "models"

st.set_page_config(
    page_title="AI Resume Job Matcher",
    page_icon="📄",
    layout="wide",
)


@st.cache_resource
def load_assets():
    required_files = [
        "logistic_regression.pkl",
        "knn.pkl",
        "tfidf_vectorizer.pkl",
        "metrics.json",
        "categories.json",
        "resume_reference.pkl",
    ]

    missing = [
        name
        for name in required_files
        if not (MODEL_DIR / name).exists()
    ]

    if missing:
        raise FileNotFoundError(
            f"Missing files: {missing}. "
            "Run python train_models.py first."
        )

    logistic = joblib.load(
        MODEL_DIR / "logistic_regression.pkl"
    )

    knn = joblib.load(
        MODEL_DIR / "knn.pkl"
    )

    vectorizer = joblib.load(
        MODEL_DIR / "tfidf_vectorizer.pkl"
    )

    with open(
        MODEL_DIR / "metrics.json",
        encoding="utf-8",
    ) as file:
        metrics = json.load(file)

    with open(
        MODEL_DIR / "categories.json",
        encoding="utf-8",
    ) as file:
        categories = json.load(file)

    reference = pd.read_pickle(
        MODEL_DIR / "resume_reference.pkl"
    )

    return (
        logistic,
        knn,
        vectorizer,
        metrics,
        categories,
        reference,
    )


def extract_resume(uploaded_file):
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(uploaded_file)

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if filename.endswith(".docx"):
        document = Document(uploaded_file)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    raise ValueError(
        "Please upload a PDF or DOCX file."
    )


st.title(
    "📄 AI Resume–Job Matching and Skill Gap Analyzer"
)

st.write(
    "Upload a resume and paste a job description to "
    "analyze similarity, missing skills, and ML predictions."
)

try:
    (
        logistic,
        knn,
        vectorizer,
        metrics,
        categories,
        reference,
    ) = load_assets()

except Exception as error:
    st.error(str(error))
    st.stop()


with st.sidebar:
    st.header("Project information")
    st.write("**ML algorithms**")
    st.write("1. Logistic Regression")
    st.write("2. K-Nearest Neighbors")
    st.write("**Text matching:** TF-IDF + cosine similarity")

    st.caption(
        "Results are informational and should not be "
        "used as the sole basis for hiring decisions."
    )


left, right = st.columns(2)

with left:
    st.subheader("1. Resume")

    uploaded_file = st.file_uploader(
        "Upload a PDF or DOCX resume",
        type=["pdf", "docx"],
    )

    pasted_resume = st.text_area(
        "Or paste resume text",
        height=230,
        placeholder="Paste resume text here...",
    )

with right:
    st.subheader("2. Job description")

    job_text = st.text_area(
        "Paste job description",
        height=300,
        placeholder=(
            "Paste the job role, requirements, "
            "responsibilities, and skills..."
        ),
    )


if st.button(
    "Analyze Resume",
    type="primary",
    use_container_width=True,
):
    resume_text = pasted_resume.strip()

    if uploaded_file is not None:
        try:
            extracted = extract_resume(
                uploaded_file
            ).strip()

            if extracted:
                resume_text = extracted

            elif not resume_text:
                st.error(
                    "No readable text was found. "
                    "This may be a scanned PDF. "
                    "Paste the text or add OCR support."
                )
                st.stop()

        except Exception as error:
            st.error(
                f"Resume extraction failed: {error}"
            )
            st.stop()

    if not resume_text:
        st.error(
            "Upload a resume or paste resume text."
        )
        st.stop()

    if not job_text.strip():
        st.error(
            "Please enter a job description."
        )
        st.stop()

    try:
        match = analyze_match(
            resume_text,
            job_text,
            vectorizer,
        )

        resume_vector = vectorizer.transform(
            [resume_text]
        )

        logistic_category = str(
            logistic.predict(resume_vector)[0]
        )

        knn_category = str(
            knn.predict(resume_vector)[0]
        )

        similar = find_similar_resumes(
            resume_text,
            vectorizer,
            knn,
            reference,
            k=5,
        )

    except Exception as error:
        st.error(
            f"Analysis failed: {error}"
        )
        st.stop()

    st.divider()
    st.header("3. Analysis results")

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Resume-job similarity",
        f"{match['similarity']}%",
    )

    coverage = match["skill_coverage"]

    col2.metric(
        "Skill coverage",
        f"{coverage}%" if coverage is not None else "N/A",
    )

    col3.metric(
        "Missing skills",
        len(match["missing_skills"]),
    )

    st.caption(
        "Similarity is a text-based indicator. "
        "Skill coverage uses only recognized dictionary skills."
    )

    st.subheader("Matched skills")

    if match["matched_skills"]:
        st.success(
            ", ".join(match["matched_skills"])
        )
    else:
        st.info("No dictionary skills matched.")

    st.subheader("Missing skills / learning priorities")

    if match["missing_skills"]:
        st.warning(
            ", ".join(match["missing_skills"])
        )

        for skill in match["missing_skills"]:
            st.write(
                f"- Consider learning or demonstrating {skill}."
            )

    elif coverage is not None:
        st.success(
            "All detected job skills appear in the resume."
        )

    else:
        st.info(
            "No recognized skills were found in the job "
            "description. Expand the dictionary in matcher.py."
        )

    st.subheader("Predicted resume category")

    col_a, col_b = st.columns(2)

    col_a.metric(
        "Logistic Regression",
        logistic_category,
    )

    col_b.metric(
        "K-NN",
        knn_category,
    )

    st.caption(
        "Category predictions are not predictions of hiring success."
    )

    st.subheader("Similar resumes in the dataset")

    if similar:
        similar_df = pd.DataFrame(
            similar
        ).rename(columns={
            "category": "Resume category",
            "similarity": "Cosine similarity (%)",
            "resume_preview": "Resume preview",
        })

        st.dataframe(
            similar_df,
            use_container_width=True,
            hide_index=True,
        )

    st.subheader("ML model comparison")

    metric_df = pd.DataFrame([
        {
            "Algorithm": name,
            "Test accuracy (%)": round(
                result["accuracy"] * 100,
                2,
            ),
        }
        for name, result in metrics.items()
    ])

    st.dataframe(
        metric_df,
        use_container_width=True,
        hide_index=True,
    )

    st.bar_chart(
        metric_df.set_index("Algorithm")
    )

    st.caption(
        "Test accuracy evaluates resume-category classification, "
        "not resume-job matching quality."
    )

    st.success("Analysis completed successfully.")
